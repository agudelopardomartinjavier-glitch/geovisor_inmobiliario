"""Exportación a Excel (.xlsx) de los reportes contables."""
from io import BytesIO

from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

def libro(hojas: dict[str, list[list]]) -> bytes:
    """`hojas`: {nombre: [fila_encabezado, fila, fila…]}. Enteros con separador de miles; encabezado en negrita; columnas ajustadas."""
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre[:31])
        for fila in filas:
            ws.append(list(fila))
        for celda in ws[1]:
            celda.font = Font(bold=True)
        for fila in ws.iter_rows(min_row=2):
            for celda in fila:
                if isinstance(celda.value, int) and not isinstance(celda.value, bool):
                    celda.number_format = "#,##0"
                elif isinstance(celda.value, float):
                    celda.number_format = "#,##0.00"
        for i, col in enumerate(ws.columns, 1):
            ancho = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, ancho + 2))
        ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

def respuesta(hojas: dict[str, list[list]], nombre: str) -> Response:
    return Response(libro(hojas), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{nombre}.xlsx"', "Cache-Control": "no-store"})
