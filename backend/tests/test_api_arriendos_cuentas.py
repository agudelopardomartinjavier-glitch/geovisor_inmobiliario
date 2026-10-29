"""Fase 3 — arriendos, cartera, recibos, gastos, cuentas del propietario, impuestos y reportes, de punta a punta sobre copias de las bases reales.
Un solo escenario con números conocidos (corte 2026-10-29)."""
import io

import pytest

from conftest import HAY_DATOS

pytestmark = pytest.mark.skipif(not HAY_DATOS, reason="faltan las bases reales")

CORTE = "2026-10-29"
PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"

def H(c):
    return c.headers_admin

def _persona(c, doc, nombre, **kw):
    r = c.post("/admin/personas", json={"nombre": nombre, "documento": doc, **kw}, headers=H(c))
    assert r.status_code == 201, r.text
    return r.json()

@pytest.fixture(scope="module")
def esc(cliente):
    """Escenario: propietario, arrendatario, codeudor, un contrato de arriendo con IPC y dos cobros vencidos."""
    h = H(cliente)
    arr = [i for i in cliente.get("/admin/inmuebles", headers=h).json() if i["operacion"] == "arriendo"]
    assert len(arr) >= 2
    prop, inq, cod = _persona(cliente, "31000001", "Propietaria Prueba"), _persona(cliente, "31000002", "Arrendatario Prueba", telefono="3001112233"), _persona(cliente, "31000003", "Codeudor Prueba")
    cliente.put("/admin/ipc/2025", json={"pct": 5.1, "fuente": "DANE"}, headers=h)
    assert cliente.put(f"/admin/inmuebles/{arr[0]['id']}/propietario", json={"persona_id": prop["id"]}, headers=h).status_code == 200
    k = cliente.post("/admin/contratos", json={"inmueble_id": arr[0]["id"], "arrendatario_id": inq["id"], "canon_inicial": 1_000_000, "fecha_inicio": "2025-03-15", "fecha_fin": "2027-03-14",
                                               "dia_pago": 5, "deposito": 1_000_000, "valor_comercial_ref": 150_000_000, "codeudores": [cod["id"]]}, headers=h)
    assert k.status_code == 201, k.text
    g = cliente.post("/admin/cobros/generar", json={"periodo": "2026-09"}, headers=h).json()
    cliente.post("/admin/cobros/generar", json={"periodo": "2026-10"}, headers=h)
    return dict(inm=arr[0], inm2=arr[1], prop=prop, inq=inq, cod=cod, k=k.json(), gen=g)

def test_rutas_exigen_sesion(cliente):
    for ruta in ("/admin/contratos", "/admin/cartera", "/admin/recibos", "/admin/gastos", "/admin/liquidaciones", "/admin/reportes/flujo-caja?desde=2026-01-01&hasta=2026-12-31",
                 "/admin/impuestos/resumen?periodo=2026-10", "/admin/contratos/alertas"):
        assert cliente.get(ruta).status_code == 401, ruta

# ---------------------------------------------------------------- contrato
def test_contrato_canon_con_ipc_deposito_y_codeudores(cliente, esc):
    k = cliente.get(f"/admin/contratos/{esc['k']['id']}?corte={CORTE}", headers=H(cliente)).json()
    assert k["canon_inicial"] == 1_000_000 and k["canon_vigente"] == 1_051_000          # aniversario 15-mar-2026 con IPC 2025 = 5,1 %
    assert [a["anio_ipc"] for a in k["aumentos"]] == [2025] and k["deposito"] == 1_000_000
    assert [x["id"] for x in k["codeudores"]] == [esc["cod"]["id"]] and k["arrendatario"]["nombre"] == "Arrendatario Prueba"
    assert k["comision_admin_pct"] == 10 and k["ley820"]["limite"] == 1_500_000 and k["ley820"]["excede"] is False

def test_validaciones_del_contrato(cliente, esc):
    h, base = H(cliente), dict(inmueble_id=esc["inm2"]["id"], arrendatario_id=esc["inq"]["id"], canon_inicial=900_000, fecha_inicio="2026-01-01", fecha_fin="2027-01-01")
    assert cliente.post("/admin/contratos", json={**base, "fecha_fin": "2025-12-31"}, headers=h).status_code == 422
    assert cliente.post("/admin/contratos", json={**base, "codeudores": [esc["inq"]["id"]]}, headers=h).status_code == 422
    assert cliente.post("/admin/contratos", json={**base, "dia_pago": 31}, headers=h).status_code == 422
    assert cliente.post("/admin/contratos", json={**base, "arrendatario_id": 999999}, headers=h).status_code == 404
    cruzado = dict(base, inmueble_id=esc["inm"]["id"], fecha_inicio="2026-06-01", fecha_fin="2027-06-01")
    assert cliente.post("/admin/contratos", json=cruzado, headers=h).status_code == 409          # se cruza con el vigente

def test_los_cobros_se_generan_una_sola_vez_con_el_canon_del_mes(cliente, esc):
    h = H(cliente)
    assert [x["valor"] for x in esc["gen"]["creados"]] == [1_051_000] and esc["gen"]["creados"][0]["fecha_vencimiento"] == "2026-09-05"
    otra = cliente.post("/admin/cobros/generar", json={"periodo": "2026-09"}, headers=h).json()
    assert otra["creados"] == [] and otra["ya_existian"] >= 1
    assert cliente.post("/admin/cobros/generar", json={"periodo": "2026-9"}, headers=h).status_code == 422
    assert [c["periodo"] for c in cliente.get(f"/admin/contratos/{esc['k']['id']}/cobros", headers=h).json()] == ["2026-09", "2026-10"]

# ---------------------------------------------------------------- cartera y mora
def test_cartera_dias_de_mora_interes_y_tramos(cliente, esc):
    r = cliente.get(f"/admin/cartera?corte={CORTE}", headers=H(cliente)).json()
    f = next(x for x in r["contratos"] if x["contrato_id"] == esc["k"]["id"])
    assert f["saldo_vencido"] == 2_102_000 and f["dias_max"] == 54 and f["tramo"] == "31-60" and f["telefono"] == "3001112233"
    d = {x["periodo"]: x for x in f["cobros"]}
    assert d["2026-09"]["dias_mora"] == 54 and d["2026-09"]["interes"] == 37_836
    assert d["2026-10"]["dias_mora"] == 24 and d["2026-10"]["interes"] == 16_816
    assert r["tramos"]["31-60"] >= 1_051_000 and r["tramos"]["1-30"] >= 1_051_000
    antes = cliente.get("/admin/cartera?corte=2026-09-04", headers=H(cliente)).json()            # antes de vencer: nada en mora de este contrato
    assert not any(x["contrato_id"] == esc["k"]["id"] for x in antes["contratos"])

def test_parametros_de_mora_cambian_el_interes(cliente, esc):
    h = H(cliente)
    cliente.put("/admin/parametros/mora.dias_gracia", json={"valor": "10"}, headers=h)
    try:
        f = next(x for x in cliente.get(f"/admin/cartera?corte={CORTE}", headers=h).json()["contratos"] if x["contrato_id"] == esc["k"]["id"])
        assert {x["periodo"]: x["dias_mora"] for x in f["cobros"]} == {"2026-09": 44, "2026-10": 14}
    finally:
        cliente.put("/admin/parametros/mora.dias_gracia", json={"valor": "0"}, headers=h)

# ---------------------------------------------------------------- recibos de caja y paz y salvo
def test_recibo_aplica_primero_lo_mas_antiguo_y_se_puede_anular(cliente, esc):
    h, kid = H(cliente), esc["k"]["id"]
    r = cliente.post("/admin/recibos", json={"valor": 1_500_000, "fecha": "2026-10-20", "medio": "transferencia", "referencia": "TRF-1", "contrato_id": kid}, headers=h)
    assert r.status_code == 201, r.text
    r = r.json()
    assert r["prefijo"] == "RC" and r["sin_aplicar"] == 0 and r["persona_id"] == esc["inq"]["id"]
    assert [(a["periodo"], a["valor"]) for a in r["aplicaciones"]] == [("2026-09", 1_051_000), ("2026-10", 449_000)]
    cob = {c["periodo"]: c for c in cliente.get(f"/admin/contratos/{kid}/cobros", headers=h).json()}
    assert cob["2026-09"]["saldo"] == 0 and cob["2026-10"]["saldo"] == 602_000
    assert cliente.get(f"/admin/recibos/{r['id']}/pdf", headers=h).content[:4] == b"%PDF"
    assert cliente.post(f"/admin/recibos/{r['id']}/anular", headers=h).json()["estado"] == "anulado"
    assert cliente.post(f"/admin/recibos/{r['id']}/anular", headers=h).status_code == 409
    assert all(c["saldo"] == 1_051_000 for c in cliente.get(f"/admin/contratos/{kid}/cobros", headers=h).json())      # vuelven a quedar pendientes

def test_recibo_con_aplicacion_manual_valida_saldos(cliente, esc):
    h, kid = H(cliente), esc["k"]["id"]
    cobros = cliente.get(f"/admin/contratos/{kid}/cobros", headers=h).json()
    oct_ = next(c for c in cobros if c["periodo"] == "2026-10")
    assert cliente.post("/admin/recibos", json={"valor": 500_000, "aplicaciones": [{"cobro_id": oct_["id"], "valor": 600_000}]}, headers=h).status_code == 422      # más que el recibo
    assert cliente.post("/admin/recibos", json={"valor": 2_000_000, "aplicaciones": [{"cobro_id": oct_["id"], "valor": 1_100_000}]}, headers=h).status_code == 422  # más que el saldo
    assert cliente.post("/admin/recibos", json={"valor": 100, "aplicaciones": [{"cobro_id": 999999, "valor": 100}]}, headers=h).status_code == 422
    ok = cliente.post("/admin/recibos", json={"valor": 1_200_000, "fecha": "2026-10-21", "aplicaciones": [{"cobro_id": oct_["id"], "valor": 1_000_000}], "persona_id": esc["inq"]["id"]}, headers=h)
    assert ok.status_code == 201 and ok.json()["sin_aplicar"] == 200_000                       # lo que sobra queda como saldo a favor
    cliente.post(f"/admin/recibos/{ok.json()['id']}/anular", headers=h)

def test_paz_y_salvo_solo_sin_saldo_vencido(cliente, esc):
    h, kid = H(cliente), esc["k"]["id"]
    assert cliente.post(f"/admin/contratos/{kid}/paz-y-salvo", json={}, headers=h).status_code == 409
    pago = cliente.post("/admin/recibos", json={"valor": 2_102_000, "fecha": "2026-10-22", "contrato_id": kid, "medio": "efectivo"}, headers=h).json()
    assert pago["sin_aplicar"] == 0
    ps = cliente.post(f"/admin/contratos/{kid}/paz-y-salvo", json={"nota": "Entrega del inmueble"}, headers=h)
    assert ps.status_code == 201 and ps.json()["numero_completo"].startswith("PS")
    pdf = cliente.get(f"/admin/paz-y-salvos/{ps.json()['id']}/pdf", headers=h)
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    assert cliente.get(f"/admin/cartera?corte={CORTE}", headers=h).json()["contratos"] == [] or not any(
        x["contrato_id"] == kid for x in cliente.get(f"/admin/cartera?corte={CORTE}", headers=h).json()["contratos"])

# ---------------------------------------------------------------- alertas
def test_alerta_de_renovacion_y_de_canon_sobre_el_limite_legal(cliente, esc):
    h = H(cliente)
    k2 = cliente.post("/admin/contratos", json={"inmueble_id": esc["inm2"]["id"], "arrendatario_id": esc["cod"]["id"], "canon_inicial": 1_500_000, "fecha_inicio": "2026-01-01",
                                                "fecha_fin": "2026-12-31", "valor_comercial_ref": 100_000_000, "incremento_ipc": False}, headers=h)
    assert k2.status_code == 201, k2.text
    a = [x for x in cliente.get(f"/admin/contratos/alertas?corte={CORTE}", headers=h).json() if x["contrato_id"] == k2.json()["id"]]
    codigos = {x["codigo"]: x for x in a}
    assert codigos["canon_ley820"]["nivel"] == "rojo" and "1.000.000" in codigos["canon_ley820"]["texto"]       # límite = 1 % de 100 M
    assert codigos["renovacion"]["nivel"] == "amarillo" and "63 días" in codigos["renovacion"]["texto"]
    lejos = cliente.get(f"/admin/contratos/{k2.json()['id']}?corte=2026-03-01", headers=h).json()
    assert "renovacion" not in {x["codigo"] for x in lejos["alertas"]}                                       # aún falta mucho
    cliente.put("/admin/parametros/canon.max_pct_valor_comercial", json={"valor": "2"}, headers=h)             # el 1 % es un parámetro
    try:
        assert "canon_ley820" not in {x["codigo"] for x in cliente.get(f"/admin/contratos/{k2.json()['id']}?corte={CORTE}", headers=h).json()["alertas"]}
    finally:
        cliente.put("/admin/parametros/canon.max_pct_valor_comercial", json={"valor": "1"}, headers=h)
    cliente.post(f"/admin/contratos/{k2.json()['id']}/terminar", json={"fecha": "2026-10-25"}, headers=h)

def test_alerta_de_ipc_faltante(cliente, esc):
    h = H(cliente)
    cliente.delete("/admin/ipc/2025", headers=h)
    try:
        k = cliente.get(f"/admin/contratos/{esc['k']['id']}?corte={CORTE}", headers=h).json()
        assert k["canon_vigente"] == 1_000_000 and k["ipc_faltante"] == [2025]
        assert "ipc_faltante" in {x["codigo"] for x in k["alertas"]}
    finally:
        cliente.put("/admin/ipc/2025", json={"pct": 5.1, "fuente": "DANE"}, headers=h)

# ---------------------------------------------------------------- gastos y soporte
def test_gastos_con_soporte_privado(cliente, esc):
    h, iid = H(cliente), esc["inm"]["id"]
    g = cliente.post("/admin/gastos", json={"inmueble_id": iid, "fecha": "2026-10-10", "categoria": "mantenimiento", "descripcion": "Cambio de grifería", "valor": 300_000, "proveedor": "Plomería SAS"}, headers=h)
    assert g.status_code == 201, g.text
    gid = g.json()["id"]
    assert g.json()["a_cargo_de"] == "propietario" and g.json()["tiene_soporte"] is False
    assert cliente.post("/admin/gastos", json={"inmueble_id": iid, "fecha": "2026-10-10", "categoria": "otro", "descripcion": "x", "valor": 0}, headers=h).status_code == 422
    assert cliente.post(f"/admin/gastos/{gid}/soporte", params={"nombre": "malo.pdf"}, content=b"MZ-ejecutable", headers=h).status_code == 415
    ok = cliente.post(f"/admin/gastos/{gid}/soporte", params={"nombre": "../factura.pdf"}, content=PDF, headers=h)
    assert ok.status_code == 200 and ok.json()["tiene_soporte"] and "soporte_archivo" not in ok.json()
    assert cliente.get(f"/admin/gastos/{gid}/soporte").status_code == 401                          # sin sesión no se descarga
    d = cliente.get(f"/admin/gastos/{gid}/soporte", headers=h)
    assert d.status_code == 200 and d.content == PDF and "attachment" in d.headers["content-disposition"]
    assert not any("factura" in str(p) and ".." in str(p) for p in (cliente.tmp / "soportes").rglob("*"))
    for categoria, subtipo, desc, valor in (("impuesto", "predial", "Predial 2026", 150_000), ("administracion_ph", None, "Cuota de administración", 100_000)):
        assert cliente.post("/admin/gastos", json={"inmueble_id": iid, "fecha": "2026-10-12", "categoria": categoria, "subtipo": subtipo, "descripcion": desc, "valor": valor}, headers=h).status_code == 201
    cliente.post("/admin/gastos", json={"inmueble_id": iid, "fecha": "2026-10-13", "categoria": "seguro", "descripcion": "Póliza (lo paga la inmobiliaria)", "valor": 80_000, "a_cargo_de": "inmobiliaria"}, headers=h)
    cliente.post("/admin/gastos", json={"inmueble_id": iid, "fecha": "2026-10-14", "categoria": "servicio", "descripcion": "Agua (arrendatario)", "valor": 70_000, "a_cargo_de": "arrendatario"}, headers=h)
    lista = cliente.get(f"/admin/gastos?inmueble_id={iid}&desde=2026-10-01&hasta=2026-10-31", headers=h).json()
    assert len(lista) == 5

# ---------------------------------------------------------------- cuenta del propietario
def test_liquidacion_mensual_con_numeros_conocidos(cliente, esc):
    h, pid = H(cliente), esc["prop"]["id"]
    v = cliente.get(f"/admin/liquidaciones/vista-previa?propietario_id={pid}&periodo=2026-10", headers=h).json()
    assert v["canones_recibidos"] == 2_102_000                                    # solo el recibo vigente del 22-oct (los anulados no cuentan)
    assert (v["comision"], v["iva_comision"]) == (210_200, 39_938)
    assert (v["administracion"], v["predial"], v["reparaciones"], v["otros"]) == (100_000, 150_000, 300_000, 0)    # lo del arrendatario y de la inmobiliaria no se descuenta
    assert v["saldo_a_girar"] == 2_102_000 - 210_200 - 39_938 - 100_000 - 150_000 - 300_000 == 1_301_862
    assert v["liquidacion_existente"] is None
    assert cliente.get(f"/admin/liquidaciones/vista-previa?propietario_id={esc['cod']['id']}&periodo=2026-10", headers=h).status_code == 422      # sin inmuebles

def test_cerrar_girar_y_pdf_de_la_liquidacion(cliente, esc):
    h, pid, iid = H(cliente), esc["prop"]["id"], esc["inm"]["id"]
    assert cliente.post("/admin/liquidaciones", json={"propietario_id": pid, "periodo": "2999-01"}, headers=h).status_code == 422       # mes futuro
    l = cliente.post("/admin/liquidaciones", json={"propietario_id": pid, "periodo": "2026-10"}, headers=h)
    assert l.status_code == 201, l.text
    l = l.json()
    assert l["numero_completo"].startswith("LQ") and l["saldo_a_girar"] == 1_301_862 and l["estado"] == "cerrada"
    assert l["detalle"]["inmuebles"][0]["recaudos"] and l["detalle"]["inmuebles"][0]["gastos"]
    assert cliente.post("/admin/liquidaciones", json={"propietario_id": pid, "periodo": "2026-10"}, headers=h).status_code == 409
    # el mes liquidado queda protegido: no se puede agregar ni borrar gastos a cargo del propietario
    assert cliente.post("/admin/gastos", json={"inmueble_id": iid, "fecha": "2026-10-30", "categoria": "otro", "descripcion": "Tarde", "valor": 5_000}, headers=h).status_code == 409
    g = cliente.get(f"/admin/gastos?inmueble_id={iid}&categoria=mantenimiento", headers=h).json()[0]
    assert cliente.delete(f"/admin/gastos/{g['id']}", headers=h).status_code == 409
    pdf = cliente.get(f"/admin/liquidaciones/{l['id']}/pdf", headers=h)
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    girada = cliente.post(f"/admin/liquidaciones/{l['id']}/girar", json={"fecha": "2026-10-28", "referencia": "TRF-777"}, headers=h).json()
    assert girada["estado"] == "girada" and girada["fecha_giro"] == "2026-10-28"
    assert cliente.post(f"/admin/liquidaciones/{l['id']}/girar", json={}, headers=h).status_code == 409
    assert cliente.delete(f"/admin/liquidaciones/{l['id']}", headers=h).status_code == 409

# ---------------------------------------------------------------- impuestos y reportes
def test_resumen_de_impuestos_y_excel(cliente, esc):
    from openpyxl import load_workbook
    h = H(cliente)
    venta = next(i for i in cliente.get("/admin/inmuebles", headers=h).json() if i["operacion"] == "venta" and i["precio"])
    com = cliente.post("/admin/comisiones", json={"inmueble_id": venta["id"], "tipo": "venta", "pagador": "comprador", "base": 200_000_000, "pct_pactado": 3, "fecha": "2026-08-10",
                                                   "asesor": "maria", "asesor_pct": 50}, headers=h).json()
    cli = _persona(cliente, "31000009", "Cliente Impuestos")
    f = cliente.post(f"/admin/comisiones/{com['id']}/facturar", json={"cliente_id": cli["id"], "fecha": "2026-08-12"}, headers=h).json()
    cliente.post(f"/admin/comisiones/{com['id']}/cobrar", json={"fecha": "2026-08-30"}, headers=h)
    cliente.post(f"/admin/facturas/{f['id']}/nota-credito", json={"motivo": "Ajuste parcial", "pct": 10, "fecha": "2026-08-20"}, headers=h)
    r = cliente.get("/admin/impuestos/resumen?periodo=2026-08", headers=h).json()
    assert (r["facturas"], r["notas_credito"]) == (1, 1)
    assert r["ingresos_netos"] == 6_000_000 - 600_000 and r["iva_generado"] == 1_140_000 - 114_000
    assert r["retefuente"] == 660_000 - 66_000 and r["ica_causado"] == round(5_400_000 * 10 / 1000) == 54_000
    assert r["ingresos_por_tipo"]["venta"] == 6_000_000 and "contador" in r["aviso"]
    x = cliente.get("/admin/impuestos/resumen?periodo=2026-08&formato=xlsx", headers=h)
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Resumen", "Facturas y notas crédito", "Comisiones"]
    resumen = {fila[0]: fila[1] for fila in wb["Resumen"].iter_rows(min_row=2, values_only=True)}
    assert resumen["Ingresos netos (antes de IVA)"] == 5_400_000 and resumen["Facturas emitidas"] == 1
    assert cliente.get("/admin/impuestos/resumen?periodo=2026-13", headers=h).status_code == 422
    assert cliente.get("/admin/cartera?corte=2026-02-30", headers=h).status_code == 422                      # fecha imposible: error del cliente, no 500
    assert cliente.post("/admin/gastos", json={"inmueble_id": 1, "fecha": "2026-02-30", "categoria": "otro", "descripcion": "x", "valor": 1}, headers=h).status_code in (409, 422)

def test_reportes_de_comisiones_flujo_de_caja_ocupacion_y_rentabilidad(cliente, esc):
    h = H(cliente)
    rc = cliente.get("/admin/reportes/comisiones?desde=2026-08-01&hasta=2026-08-31", headers=h).json()
    assert rc["total_cobrado"] >= 6_000_000 and any(m["mes"] == "2026-08" for m in rc["por_mes"]) and rc["por_tipo"]["venta"] >= 6_000_000
    assert any(a["nombre"] == "maria" and a["valor"] >= 3_000_000 for a in rc["por_asesor"])
    fc = cliente.get("/admin/reportes/flujo-caja?desde=2026-10-01&hasta=2026-10-31", headers=h).json()
    oct_ = next(m for m in fc["meses"] if m["mes"] == "2026-10")
    assert oct_["recaudos"] == 2_102_000                                           # recibos vigentes de octubre
    assert oct_["gastos_pagados"] == 300_000 + 150_000 + 100_000 + 80_000          # propietario + inmobiliaria; el del arrendatario no sale de nuestra caja
    assert oct_["giros"] == 1_301_862 and oct_["neto"] == oct_["entradas"] - oct_["salidas"]
    oc = cliente.get(f"/admin/reportes/ocupacion?corte={CORTE}", headers=h).json()
    assert oc["ocupados"] >= 1 and oc["inmuebles_administrados"] == oc["ocupados"] + oc["vacantes"] and 0 <= oc["ocupacion_pct"] <= 100
    assert any(o["inmueble_id"] == esc["inm"]["id"] for o in oc["detalle_ocupados"])
    rt = cliente.get(f"/admin/reportes/rentabilidad?corte={CORTE}&meses=12", headers=h).json()
    f = next(x for x in rt["inmuebles"] if x["inmueble_id"] == esc["inm"]["id"])
    assert f["canon_mensual"] == 1_051_000 and f["valor"] == 150_000_000
    assert f["rentabilidad_teorica_pct"] == round(100 * 1_051_000 * 12 / 150_000_000, 2)
    assert f["ingresos_periodo"] == 2_102_000 and f["gastos_propietario_periodo"] == 550_000
    assert f["rentabilidad_real_pct"] == round(100 * (2_102_000 - 550_000) * 1 / 150_000_000, 2)
    for ruta in ("reportes/comisiones?desde=2026-01-01&hasta=2026-12-31", "reportes/flujo-caja?desde=2026-01-01&hasta=2026-12-31", f"reportes/ocupacion?corte={CORTE}", f"reportes/rentabilidad?corte={CORTE}"):
        sep = "&" if "?" in ruta else "?"
        assert cliente.get(f"/admin/{ruta}{sep}formato=xlsx", headers=h).status_code == 200

def test_un_inmueble_con_movimientos_contables_no_se_borra(cliente, esc):
    assert cliente.delete(f"/admin/inmuebles/{esc['inm']['id']}", headers=H(cliente)).status_code == 409

def test_la_auditoria_contable_no_guarda_datos_personales(cliente, esc):
    import sqlite3
    con = sqlite3.connect(cliente.tmp / "admin.sqlite")
    texto = " ".join(str(r) for r in con.execute("SELECT accion, detalle FROM auditoria WHERE accion LIKE 'persona%' OR accion LIKE 'recibo%' OR accion LIKE 'contrato%' OR accion LIKE 'liquidacion%'"))
    con.close()
    assert "contrato_crear" in texto and "liquidacion_cerrar" in texto
    for dato in ("Arrendatario Prueba", "Propietaria Prueba", "31000002", "3001112233"):
        assert dato not in texto
